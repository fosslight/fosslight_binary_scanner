#!/usr/bin/env python
# -*- coding: utf-8 -*-
# Copyright (c) 2026 LG Electronics Inc.
# SPDX-License-Identifier: Apache-2.0

import hashlib
from copy import deepcopy
from random import Random

import pytest
import yaml
from openpyxl import load_workbook
from spdx_tools.spdx.model import FileType, RelationshipType
from spdx_tools.spdx.parser.parse_anything import parse_file
from spdx_tools.spdx.validation.document_validator import validate_full_spdx_document

from fosslight_binary import binary_analysis
from fosslight_util import write_spdx as spdx_writer
from fosslight_util.oss_item import OssItem


@pytest.fixture
def binary_project(tmp_path, monkeypatch):
    project = tmp_path / "project"
    project.mkdir()
    binary = project / "artifact.bin"
    binary.write_bytes(Random(4519).randbytes(8192))
    output = tmp_path / "reports"

    def lookup(items, _url, _token):
        for item in items:
            item.set_oss_items([OssItem("first", "1.0", "MIT"), OssItem("second", "2.0", "Apache-2.0")])
        return items, len(items), ""

    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(binary_analysis, "get_oss_info_from_db", lookup)
    return project, binary, output


@pytest.mark.parametrize("format_name, extension", [
    ("spdx-json", ".json"), ("spdx-yaml", ".yaml"), ("spdx-xml", ".xml"), ("spdx-tag", ".tag"),
])
@pytest.mark.parametrize("spdx_first", [False, True])
def test_binary_scan_preserves_hashes_and_shared_file(binary_project, monkeypatch, format_name, extension, spdx_first):
    project, binary, output = binary_project

    def unexpected_checksum_lookup(_filename):
        raise AssertionError("The Binary scanner SHA-1 must be reused")

    monkeypatch.setattr(spdx_writer, "get_checksum_sha1", unexpected_checksum_lookup)
    formats = [format_name, "excel"] if spdx_first else ["excel", format_name]
    success, scan_item = binary_analysis.find_binaries(str(project), str(output), formats, correct_mode=False)

    assert success
    item, = scan_item.file_items[binary_analysis.PKG_NAME]
    assert item.source_name_or_path == "artifact.bin"
    assert item.checksum == hashlib.sha1(binary.read_bytes()).hexdigest()
    assert item.tlsh not in ("", "0", "TNULL")
    filename, = output.glob(f"*{extension}")
    document = parse_file(str(filename))
    assert validate_full_spdx_document(document) == []
    assert document.creation_info.name == "project"
    assert any(creator.name.startswith("FOSSLIGHT_Binary - ") for creator in document.creation_info.creators)
    assert len(document.packages) == 3
    file_info, = document.files
    assert file_info.name == "./artifact.bin"
    assert file_info.file_types == [FileType.BINARY]
    assert file_info.checksums[0].value == item.checksum
    assert file_info.comment == f"fosslight:tlsh={item.tlsh}"
    assert {str(value) for value in file_info.license_info_in_file} == {"MIT", "Apache-2.0"}
    packages = document.packages[1:]
    for package in packages:
        assert package.files_analyzed is True
        assert package.license_declared is None
        assert {str(value) for value in package.license_info_from_files} == {"MIT", "Apache-2.0"}
    assert {
        (relation.spdx_element_id, relation.related_spdx_element_id)
        for relation in document.relationships if relation.relationship_type == RelationshipType.CONTAINS
    } == {(package.spdx_id, file_info.spdx_id) for package in packages}
    filename, = output.glob("*.xlsx")
    workbook = load_workbook(filename, read_only=True)
    try:
        assert workbook["BIN_FL_Binary"].max_row == 3
    finally:
        workbook.close()


def test_binary_correction_uses_returned_scanner_item(binary_project, monkeypatch):
    project, _binary, output = binary_project

    def correct(_correct_path, _scan_path, scan_item):
        corrected = deepcopy(scan_item)
        item, = corrected.file_items[binary_analysis.PKG_NAME]
        item.oss_items[0].name = "corrected"
        item.oss_items[1].exclude = True
        return True, "", corrected

    monkeypatch.setattr(binary_analysis, "correct_with_yaml", correct)
    success, scan_item = binary_analysis.find_binaries(str(project), str(output), ["spdx-json"])

    assert success
    assert scan_item.file_items[binary_analysis.PKG_NAME][0].oss_items[0].name == "corrected"
    filename, = output.glob("*.json")
    document = parse_file(str(filename))
    assert validate_full_spdx_document(document) == []
    assert [package.name for package in document.packages] == ["project", "corrected"]
    assert {str(value) for value in document.files[0].license_info_in_file} == {"MIT"}


@pytest.mark.parametrize("format_name, extension", [
    ("spdx-json", ".json"), ("spdx-yaml", ".yaml"), ("spdx-xml", ".xml"), ("spdx-tag", ".tag"),
])
def test_real_yaml_correction_preserves_binary_hashes(binary_project, format_name, extension):
    project, binary, output = binary_project
    correction = {"corrected-package": [{"version": "3.0", "license": "BSD-3-Clause", "binary path": "artifact.bin"}]}
    (project / "sbom-info.yaml").write_text(yaml.safe_dump(correction), encoding="utf-8")

    success, scan_item = binary_analysis.find_binaries(str(project), str(output), [format_name])

    assert success
    items = scan_item.file_items[binary_analysis.PKG_NAME]
    assert len(items) == 2
    assert items[0].exclude is True
    filename, = output.glob(f"*{extension}")
    document = parse_file(str(filename))
    assert validate_full_spdx_document(document) == []
    assert [package.name for package in document.packages] == ["project", "corrected-package"]
    file_info, = document.files
    assert file_info.checksums[0].value == hashlib.sha1(binary.read_bytes()).hexdigest()
    assert file_info.comment == f"fosslight:tlsh={items[1].tlsh}"
    assert {str(value) for value in file_info.license_info_in_file} == {"BSD-3-Clause"}


@pytest.mark.parametrize("exclusion", ["file", "all_oss", "one_oss"])
def test_binary_exclusions_apply_to_files_packages_and_licenses(binary_project, monkeypatch, exclusion):
    project, _binary, output = binary_project
    original_lookup = binary_analysis.get_oss_info_from_db

    def lookup(items, url, token):
        items, count, comment = original_lookup(items, url, token)
        item, = items
        if exclusion == "file":
            item.exclude = True
        else:
            item.oss_items[1].exclude = True
            if exclusion == "all_oss":
                item.oss_items[0].exclude = True
        return items, count, comment

    monkeypatch.setattr(binary_analysis, "get_oss_info_from_db", lookup)
    success, _scan_item = binary_analysis.find_binaries(str(project), str(output), ["spdx-json"], correct_mode=False)

    assert success
    filename, = output.glob("*.json")
    document = parse_file(str(filename))
    assert validate_full_spdx_document(document) == []
    assert len(document.files) == (1 if exclusion == "one_oss" else 0)
    assert len(document.packages) == (2 if exclusion == "one_oss" else 1)
    if exclusion == "one_oss":
        assert {str(value) for value in document.files[0].license_info_in_file} == {"MIT"}


@pytest.mark.parametrize("empty", [False, True])
def test_binary_unmatched_and_empty_results_keep_valid_document(binary_project, monkeypatch, empty):
    project, binary, output = binary_project
    if empty:
        binary.unlink()
    monkeypatch.setattr(binary_analysis, "get_oss_info_from_db", lambda items, *_args: (items, 0, ""))

    success, _scan_item = binary_analysis.find_binaries(str(project), str(output), ["spdx-json"], correct_mode=False)

    assert success
    filename, = output.glob("*.json")
    document = parse_file(str(filename))
    assert validate_full_spdx_document(document) == []
    assert len(document.packages) == 1
    assert len(document.files) == (0 if empty else 1)
    assert any(creator.name.startswith("FOSSLIGHT_Binary - ") for creator in document.creation_info.creators)


def test_binary_without_valid_tlsh_keeps_sha1_but_omits_comment(binary_project):
    project, binary, output = binary_project
    binary.write_bytes(b"\x00small binary\x01")

    success, scan_item = binary_analysis.find_binaries(str(project), str(output), ["spdx-json"], correct_mode=False)

    assert success
    assert scan_item.file_items[binary_analysis.PKG_NAME][0].tlsh == "0"
    filename, = output.glob("*.json")
    document = parse_file(str(filename))
    assert validate_full_spdx_document(document) == []
    assert document.files[0].comment is None
    assert document.files[0].checksums[0].value == hashlib.sha1(binary.read_bytes()).hexdigest()


def test_yaml_correction_keeps_existing_report_formats(binary_project):
    project, _binary, output = binary_project
    correction = {"corrected-package": [{"version": "3.0", "license": "BSD-3-Clause", "binary path": "artifact.bin"}]}
    (project / "sbom-info.yaml").write_text(yaml.safe_dump(correction), encoding="utf-8")

    success, _scan_item = binary_analysis.find_binaries(
        str(project), str(output), ["excel", "csv", "yaml", "opossum", "spdx-json"],
    )

    assert success
    filename, = output.glob("*.xlsx")
    workbook = load_workbook(filename, read_only=True)
    try:
        assert workbook["BIN_FL_Binary"].max_row == 4
        assert any(row[2] == "corrected-package" for row in workbook["BIN_FL_Binary"].iter_rows(min_row=2, values_only=True))
    finally:
        workbook.close()
    for pattern in ("*.csv", "*.yaml", "*opossum*.json"):
        filename, = output.glob(pattern)
        assert "corrected-package" in filename.read_text(encoding="utf-8")
